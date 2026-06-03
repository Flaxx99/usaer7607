"""
test_exports.py — Validación de exportaciones Excel.

Verifica que los archivos Excel generados contengan datos reales.
"""

import openpyxl
from io import BytesIO

from django.urls import reverse
from rest_framework import status

from .base import BaseIntegrationTest
from rac.models import RegistroRAC


class ExportTest(BaseIntegrationTest):

    def test_export_rac_excel_tiene_datos(self):
        """Admin exporta RAC y el Excel contiene datos del alumno."""
        RegistroRAC.objects.create(
            alumno=self.alumno,
            ciclo_escolar=self.ciclo,
            escuela_regular=self.escuela,
            escuela_basica=self.escuela,
            maestro_apoyo=self.maestro,
            clasificacion="DISCAPACIDAD",
            subclasificacion="DI",
            service_type="USAER",
        )

        self._auth(self.admin)
        url = reverse("rac:exportar_todo")
        response = self.client.get(url)
        self.assertEqual(response.status_code, status.HTTP_200_OK)

        wb = openpyxl.load_workbook(BytesIO(response.content))
        ws = wb["RAC"]
        found = False
        for row in ws.iter_rows(values_only=True):
            if any(self.alumno.apellido_paterno in str(cell or "")
                   for cell in row):
                found = True
                break
        self.assertTrue(found, "El Excel exportado debe contener datos del alumno")

    def test_export_rae_excel_tiene_datos(self):
        """Admin exporta RAE individual y el Excel contiene datos."""
        self._auth(self.maestro)
        init_url = reverse("rae:captura_rae")
        init_resp = self.client.get(init_url)
        registro_id = init_resp.json()["registro_id"]

        self._auth(self.admin)
        export_url = reverse("rae:exportar_rae_excel", args=[registro_id])
        response = self.client.get(export_url)
        self.assertEqual(response.status_code, status.HTTP_200_OK)

        # RAE usa get_full_name() -> "NOMBRES APELLIDO_PATERNO APELLIDO_MATERNO"
        nombre_completo = self.alumno.get_full_name()

        wb = openpyxl.load_workbook(BytesIO(response.content))
        ws = wb["Sheet1"]
        found = False
        for row in ws.iter_rows(values_only=True):
            if any(nombre_completo in str(cell or "").upper()
                   for cell in row):
                found = True
                break
        self.assertTrue(found, "El Excel exportado debe contener el nombre completo del alumno")
