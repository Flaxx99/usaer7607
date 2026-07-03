# usaer_system/rae/views.py
import logging
import os
from io import BytesIO

from alumnos.models import Alumno
from ciclos_escolares.utils import get_current_ciclo_escolar_instance
from django.conf import settings
from django.db import transaction
from django.db.models import Q
from django.http import HttpResponse
from django.shortcuts import get_object_or_404
from openpyxl import load_workbook
from rest_framework import filters, permissions, views, viewsets
from rest_framework.response import Response
from services.dto import (
    CerrarRegistroResponse,
    RAEAlumnoResponse,
    RAEInitResponse,
    RAEProgressItem,
    StatusDetailResponse,
    ToggleCerradoPayload,
)
from services.error_handling import error_400, error_403, error_500
from usuarios.models import SystemConfiguration

from .models import RAEAlumno, RegistroRAE
from .permissions import RAEPermission
from .serializers import BulkRAESaveSerializer, RAEAlumnoSerializer, RegistroRAESerializer

logger = logging.getLogger(__name__)


class RegistroRAEViewSet(viewsets.ModelViewSet):
    serializer_class = RegistroRAESerializer
    permission_classes = [permissions.IsAuthenticated, RAEPermission]
    filter_backends = [filters.SearchFilter, filters.OrderingFilter]
    search_fields = ["escuela__nombre"]
    ordering_fields = ["fecha_creacion"]

    def get_queryset(self):
        user = self.request.user
        qs = RegistroRAE.objects.select_related("escuela", "ciclo_escolar", "creado_por")
        try:
            ciclo_actual = get_current_ciclo_escolar_instance()
            qs = qs.filter(ciclo_escolar=ciclo_actual)
        except Exception:
            return RegistroRAE.objects.none()
        roles_totales = ["ADMIN", "SECRETARIO"]
        if not (user.is_superuser or getattr(user, "role", "") in roles_totales):
            if hasattr(user, "escuela") and user.escuela:
                qs = qs.filter(escuela=user.escuela)
            else:
                return RegistroRAE.objects.none()
        return qs.order_by("escuela__nombre", "fecha_creacion")


class RAEInitCaptureView(views.APIView):
    """Inicializa la captura RAE para el usuario autenticado.

    Crea el registro del ciclo si no existe, instancia los RAEAlumno
    para los alumnos del docente, y devuelve la versión actual para
    bloqueo optimista.
    """

    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        user = request.user
        escuela = user.escuela
        if not escuela:
            return error_400("No tienes escuela asignada.")
        try:
            ciclo = get_current_ciclo_escolar_instance()
        except Exception as e:
            return error_400(str(e))

        registro, _ = RegistroRAE.objects.get_or_create(
            escuela=escuela,
            ciclo_escolar=ciclo,
            defaults={"creado_por": user},
        )
        alumnos_maestra = Alumno.objects.activos_de_profesor(user)
        rae_alumnos = []
        for alumno in alumnos_maestra:
            obj, _ = RAEAlumno.objects.get_or_create(
                registro=registro,
                alumno=alumno,
                defaults={
                    "capturado_por": user,
                    "curp": alumno.curp,
                    "genero": alumno.sexo,
                    "edad": alumno.edad,
                    "grado": f"{alumno.grado}°{alumno.grupo}",
                },
            )
            rae_alumnos.append(obj)

        serializer = RAEAlumnoSerializer(rae_alumnos, many=True)
        return Response(
            RAEInitResponse(
                registro_id=registro.id,
                version=registro.version,
                ciclo=ciclo.nombre,
                escuela=escuela.nombre,
                cerrado=registro.cerrado,
                alumnos=[RAEAlumnoResponse(**a) for a in serializer.data],
            ).model_dump(),
        )


class RAEBulkSaveView(views.APIView):
    """Guarda cambios masivos de RAE con bloqueo optimista.

    El cliente envía la `version` que obtuvo al inicializar.
    Si otro usuario ya modificó el registro, devuelve 409 Conflict.
    Si coincide, aplica los cambios e incrementa la versión.
    """

    permission_classes = [permissions.IsAuthenticated]

    def post(self, request):
        serializer = BulkRAESaveSerializer(data=request.data)
        if not serializer.is_valid():
            return Response(serializer.errors, status=400)

        data = serializer.validated_data
        reg_id = data["registro_id"]
        client_ver = data["version"]

        # Filtro por escuela si no es admin
        qs_reg = RegistroRAE.objects.filter(pk=reg_id)
        if not (request.user.is_superuser or request.user.role in ["ADMIN", "SECRETARIO"]):
            qs_reg = qs_reg.filter(escuela=request.user.escuela)
        registro = get_object_or_404(qs_reg)

        # ── Bloqueo optimista ──────────────────────────────────────────────
        if registro.version != client_ver:
            return Response(
                {
                    "error": "Conflicto",
                    "detail": "El registro fue modificado por otro usuario. Recargá la página y volvé a intentar.",
                },
                status=409,
            )

        if registro.cerrado:
            return error_403("Registro cerrado. No se pueden modificar datos.")

        campos = [
            "ceg",
            "bv",
            "so",
            "hp",
            "scg",
            "dmo",
            "di",
            "dme",
            "dm",
            "psicosocial",
            "dsc",
            "dsco",
            "dsa",
            "tea",
            "tda",
            "asi",
            "asc",
            "ass",
            "asa",
            "asp",
            "ot",
            "psicologia",
            "comunicacion",
            "psicomotricidad",
            "trabajo_social",
            "aprendizaje",
            "nuevo_ingreso",
            "subsecuente",
            "diagnostico",
            "educativo",
            "deteccion",
            "psicopedagogico",
            "plan",
            "modelo",
        ]

        updated_count = 0
        instances = []

        with transaction.atomic():
            for alum_data in data["alumnos"]:
                pk = alum_data.get("id")
                try:
                    inst = RAEAlumno.objects.get(pk=pk, registro=registro)
                except RAEAlumno.DoesNotExist:
                    continue

                changed = False
                for f in campos:
                    if f in alum_data and getattr(inst, f) != bool(alum_data[f]):
                        setattr(inst, f, bool(alum_data[f]))
                        changed = True

                if changed:
                    instances.append(inst)
                    updated_count += 1

            if instances:
                RAEAlumno.objects.bulk_update(instances, fields=campos)

            # Incrementar versión del registro
            registro.version += 1
            registro.save(update_fields=["version"])

        return Response(
            StatusDetailResponse(detail=f"{updated_count} alumnos actualizados").model_dump(),
        )


class ExportRAEView(views.APIView):
    """Exporta un registro RAE a Excel usando la plantilla oficial."""

    permission_classes = [permissions.IsAuthenticated]

    def get(self, request, pk):
        qs = RegistroRAE.objects.select_related("escuela", "ciclo_escolar")
        if not (request.user.is_superuser or request.user.role in ["ADMIN", "SECRETARIO"]):
            qs = qs.filter(escuela=request.user.escuela)
        registro = get_object_or_404(qs, pk=pk)
        config = SystemConfiguration.objects.first()
        template_path = os.path.join(
            settings.BASE_DIR, "rae", "static", "excel_templates", "rae_template.xlsx"
        )
        if not os.path.exists(template_path):
            return error_500("Plantilla no encontrada.")
        try:
            wb = load_workbook(template_path)
            ws = wb.active or wb["Sheet1"]
            ws["D6"] = registro.escuela.nombre
            ws["AC6"] = (
                registro.escuela.get_nivel_display()
                if hasattr(registro.escuela, "get_nivel_display")
                else registro.escuela.nivel
            )
            ws["D8"] = registro.escuela.zona or ""
            ws["H8"] = registro.escuela.clave_estatal or ""
            ws["R8"] = registro.escuela.cct
            ws["AC8"] = registro.ciclo_escolar.nombre
            ws["D10"] = registro.escuela.domicilio or ""
            ws["D12"] = config.centro_nombre if config else "USAER 7607"
            ws["R12"] = config.centro_cct if config else "08FUA0093E"
            ws["D14"] = config.director_responsable if config else "S/N"
            ws["C80"] = ws["D14"].value
            ws["AG14"] = registro.docente_hombres
            ws["AO14"] = registro.docente_mujeres
            roles_full = ["ADMIN", "SECRETARIO"]
            qs_rae = RAEAlumno.objects.filter(registro=registro)
            if not (request.user.is_superuser or getattr(request.user, "role", "") in roles_full):
                qs_rae = qs_rae.filter(alumno__profesor=request.user)
            qs_rae = qs_rae.order_by(
                "alumno__grado", "alumno__grupo", "alumno__apellido_paterno", "alumno__nombres"
            ).select_related("alumno__profesor")
            apt_q = Q()
            for f in ["asi", "asc", "ass", "asa", "asp"]:
                apt_q |= Q(**{f: True})
            ws["AE11"] = qs_rae.filter(apt_q, genero="H").count()
            ws["AF11"] = qs_rae.filter(apt_q, genero="M").count()
            ws["AG11"] = (ws["AE11"].value or 0) + (ws["AF11"].value or 0)
            disc_q = Q()
            for f in ["ceg", "bv", "so", "hp", "scg", "dmo", "di", "dme", "psicosocial", "dm"]:
                disc_q |= Q(**{f: True})
            ws["AJ11"] = qs_rae.filter(disc_q, genero="H").count()
            ws["AK11"] = qs_rae.filter(disc_q, genero="M").count()
            ws["AL11"] = (ws["AJ11"].value or 0) + (ws["AK11"].value or 0)
            otras_q = Q()
            for f in ["ot", "dsc", "dsco", "dsa", "tea", "tda"]:
                otras_q |= Q(**{f: True})
            ws["AP11"] = qs_rae.filter(otras_q, genero="H").count()
            ws["AQ11"] = qs_rae.filter(otras_q, genero="M").count()
            ws["AR11"] = (ws["AP11"].value or 0) + (ws["AQ11"].value or 0)
            for r in range(19, 69):
                for c in range(1, 45):
                    if c != 2:
                        ws.cell(row=r, column=c).value = None
            col_map = {
                "ceg": "H",
                "bv": "I",
                "so": "J",
                "hp": "K",
                "scg": "L",
                "dmo": "M",
                "di": "N",
                "dme": "O",
                "psicosocial": "P",
                "dm": "Q",
                "dsc": "R",
                "dsco": "S",
                "dsa": "T",
                "tea": "U",
                "tda": "V",
                "asi": "W",
                "asc": "X",
                "ass": "Y",
                "asa": "Z",
                "asp": "AA",
                "ot": "AB",
                "psicologia": "AC",
                "comunicacion": "AD",
                "psicomotricidad": "AE",
                "trabajo_social": "AF",
                "aprendizaje": "AG",
                "nuevo_ingreso": "AI",
                "subsecuente": "AJ",
                "diagnostico": "AL",
                "educativo": "AM",
                "deteccion": "AN",
                "psicopedagogico": "AO",
                "plan": "AP",
                "modelo": "AQ",
            }
            curr_row = 19
            for idx, rae in enumerate(qs_rae, 1):
                ws[f"B{curr_row}"] = idx
                ws[f"C{curr_row}"] = rae.alumno.get_full_name()
                ws[f"D{curr_row}"] = rae.genero or ""
                ws[f"E{curr_row}"] = rae.edad or ""
                ws[f"F{curr_row}"] = rae.grado or ""
                ws[f"G{curr_row}"] = rae.curp or ""
                for field, col in col_map.items():
                    if getattr(rae, field):
                        ws[f"{col}{curr_row}"] = "X"
                if rae.alumno.profesor:
                    ws[f"AR{curr_row}"] = (
                        rae.alumno.profesor.get_full_name() or rae.alumno.profesor.nombre
                    )
                curr_row += 1
            output = BytesIO()
            wb.save(output)
            output.seek(0)
            filename = (
                f"RAE_{registro.escuela.nombre}_{registro.ciclo_escolar.nombre}.xlsx".replace(
                    " ", "_"
                )
            )
            response = HttpResponse(
                output,
                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
            response["Content-Disposition"] = f'attachment; filename="{filename}"'
            return response
        except Exception as e:
            logger.exception("Error al exportar RAE")
            return error_500(f"Error al generar el Excel: {e}")


class ExportAllRAEView(views.APIView):
    """Exporta todos los registros RAE del ciclo actual a Excel."""

    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

        try:
            ciclo = get_current_ciclo_escolar_instance()
        except Exception as e:
            return error_400(str(e))
        qs = RegistroRAE.objects.filter(ciclo_escolar=ciclo).select_related("escuela")
        if not (request.user.is_superuser or request.user.role in ["ADMIN", "SECRETARIO"]):
            qs = qs.filter(escuela=request.user.escuela)
        wb = Workbook()
        ws = wb.active
        ws.title = "RAE General"
        header_fill = PatternFill(start_color="1e40af", end_color="1e40af", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True, size=11)
        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )
        headers = [
            "Escuela",
            "CCT",
            "Ciclo",
            "Docentes H",
            "Docentes M",
            "Total Alumnos",
            "Aptitudes",
            "Discapacidades",
            "Otras Condiciones",
            "Cerrado",
        ]
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=h)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
        ws.column_dimensions["A"].width = 30
        ws.column_dimensions["B"].width = 18
        ws.column_dimensions["C"].width = 14
        for col_letter in ["D", "E", "F", "G", "H", "I", "J", "K"]:
            ws.column_dimensions[col_letter].width = 14
        row = 2
        for reg in qs:
            raes = RAEAlumno.objects.filter(registro=reg)
            total = raes.count()
            apt_q = Q()
            for f in ["asi", "asc", "ass", "asa", "asp"]:
                apt_q |= Q(**{f: True})
            disc_q = Q()
            for f in ["ceg", "bv", "so", "hp", "scg", "dmo", "di", "dme", "psicosocial", "dm"]:
                disc_q |= Q(**{f: True})
            otras_q = Q()
            for f in ["ot", "dsco", "dsa", "dsc", "tea", "tda"]:
                otras_q |= Q(**{f: True})
            ws.cell(row=row, column=1, value=reg.escuela.nombre).border = thin_border
            ws.cell(row=row, column=2, value=reg.escuela.cct).border = thin_border
            ws.cell(row=row, column=3, value=reg.ciclo_escolar.nombre).border = thin_border
            ws.cell(row=row, column=4, value=reg.docente_hombres).border = thin_border
            ws.cell(row=row, column=5, value=reg.docente_mujeres).border = thin_border
            ws.cell(row=row, column=6, value=total).border = thin_border
            ws.cell(row=row, column=7, value=raes.filter(apt_q).count()).border = thin_border
            ws.cell(row=row, column=8, value=raes.filter(disc_q).count()).border = thin_border
            ws.cell(row=row, column=9, value=raes.filter(otras_q).count()).border = thin_border
            ws.cell(row=row, column=10, value="Sí" if reg.cerrado else "No").border = thin_border
            row += 1
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        filename = f"RAE_Todos_{ciclo.nombre}.xlsx".replace(" ", "_")
        response = HttpResponse(
            output, content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        return response


class RAECerrarView(views.APIView):
    """Cierra o reabre un registro RAE."""

    permission_classes = [permissions.IsAuthenticated, RAEPermission]

    def post(self, request, pk):
        qs = RegistroRAE.objects.filter(pk=pk)
        if not (request.user.is_superuser or request.user.role in ["ADMIN", "SECRETARIO"]):
            qs = qs.filter(escuela=request.user.escuela)
        registro = get_object_or_404(qs)
        serializer = ToggleCerradoPayload(**request.data)
        registro.cerrado = serializer.cerrado
        registro.save(update_fields=["cerrado"])
        return Response(
            CerrarRegistroResponse(
                detail="Registro cerrado." if serializer.cerrado else "Registro reabierto.",
                registro_id=registro.id,
                cerrado=registro.cerrado,
            ).model_dump(),
        )


class RAEProgressView(views.APIView):
    """Devuelve el progreso de captura RAE por escuela."""

    permission_classes = [permissions.IsAuthenticated]

    def get(self, request):
        try:
            ciclo = get_current_ciclo_escolar_instance()
        except Exception as e:
            return error_400(str(e))
        qs = RegistroRAE.objects.filter(ciclo_escolar=ciclo).select_related("escuela")
        if not (request.user.is_superuser or request.user.role in ["ADMIN", "SECRETARIO"]):
            qs = qs.filter(escuela=request.user.escuela)
        items = []
        for reg in qs:
            total = RAEAlumno.objects.filter(registro=reg).count()
            completados = (
                RAEAlumno.objects.filter(registro=reg)
                .exclude(
                    ceg=False,
                    bv=False,
                    so=False,
                    hp=False,
                    scg=False,
                    dmo=False,
                    di=False,
                    dme=False,
                    psicosocial=False,
                    dm=False,
                    dsc=False,
                    dsco=False,
                    dsa=False,
                    tea=False,
                    tda=False,
                    asi=False,
                    asc=False,
                    ass=False,
                    asa=False,
                    asp=False,
                    ot=False,
                )
                .count()
            )
            items.append(
                RAEProgressItem(
                    escuela_id=reg.escuela.id,
                    escuela_nombre=reg.escuela.nombre,
                    escuela_cct=reg.escuela.cct,
                    registro_id=reg.id,
                    total_alumnos=total,
                    completados=completados,
                    porcentaje=round((completados / total * 100) if total else 0, 1),
                    cerrado=reg.cerrado,
                ).model_dump(),
            )
        return Response(items)
